"""数据导出服务."""

import csv
import json
from pathlib import Path

import structlog
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from foshan_scraper.models import LandTransaction

log = structlog.get_logger()

# Excel 列标题映射
EXCEL_HEADERS = {
    "row_guid": "记录ID",
    "transaction_date": "成交时间",
    "winner": "竞得人",
    "transaction_area": "交易土地面积",
    "transaction_price": "成交地价",
    "land_location": "地块位置",
    "land_usage": "土地实际用途",
    "detail_url": "详情链接",
    "crawled_at": "爬取时间",
    "land_code": "地块编号",
    "supply_type": "供应类型",
    "trade_type": "交易方式",
    "start_price": "起始价",
    "status": "交易状态",
}

# 列顺序
EXCEL_FIELDS = list(EXCEL_HEADERS.keys())


def export_to_json(transactions: list[LandTransaction], output_path: Path) -> None:
    """导出数据到 JSON 文件.

    Args:
        transactions: 交易记录列表
        output_path: 输出文件路径
    """
    data = [t.model_dump(mode="json") for t in transactions]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2, default=str)

    log.info("json_export_completed", path=str(output_path), count=len(transactions))


def export_to_csv(transactions: list[LandTransaction], output_path: Path) -> None:
    """导出数据到 CSV 文件.

    Args:
        transactions: 交易记录列表
        output_path: 输出文件路径
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=EXCEL_FIELDS, extrasaction="ignore")
        writer.writeheader()
        for t in transactions:
            row = t.model_dump()
            writer.writerow(row)

    log.info("csv_export_completed", path=str(output_path), count=len(transactions))


def export_to_excel(transactions: list[LandTransaction], output_path: Path) -> None:
    """导出数据到 Excel 文件.

    Args:
        transactions: 交易记录列表
        output_path: 输出文件路径
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    if ws is None:
        msg = "Failed to create worksheet"
        raise RuntimeError(msg)
    ws.title = "土地交易记录"

    # 写入表头
    for col_idx, field in enumerate(EXCEL_FIELDS, 1):
        ws.cell(row=1, column=col_idx, value=EXCEL_HEADERS[field])

    # 写入数据
    for row_idx, t in enumerate(transactions, 2):
        row_data = t.model_dump()
        for col_idx, field in enumerate(EXCEL_FIELDS, 1):
            value = row_data.get(field, "")
            ws.cell(row=row_idx, column=col_idx, value=value)

    # 自动调整列宽
    for col_idx in range(1, len(EXCEL_FIELDS) + 1):
        col_letter = get_column_letter(col_idx)
        max_length = len(EXCEL_HEADERS[EXCEL_FIELDS[col_idx - 1]])
        for row_idx in range(2, len(transactions) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value:
                max_length = max(max_length, len(str(cell_value)))
        ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

    wb.save(output_path)
    log.info("excel_export_completed", path=str(output_path), count=len(transactions))


def export_all(
    transactions: list[LandTransaction],
    output_dir: Path,
    prefix: str = "land_transactions",
) -> tuple[Path, Path, Path]:
    """同时导出 JSON、CSV 和 Excel.

    Args:
        transactions: 交易记录列表
        output_dir: 输出目录
        prefix: 文件名前缀

    Returns:
        (json_path, csv_path, excel_path) 元组
    """
    json_path = output_dir / f"{prefix}.json"
    csv_path = output_dir / f"{prefix}.csv"
    excel_path = output_dir / f"{prefix}.xlsx"

    export_to_json(transactions, json_path)
    export_to_csv(transactions, csv_path)
    export_to_excel(transactions, excel_path)

    return json_path, csv_path, excel_path
